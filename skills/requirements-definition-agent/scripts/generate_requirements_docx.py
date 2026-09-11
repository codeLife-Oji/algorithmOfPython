"""要件定義JSONから要件定義書（Word）を生成する。

使い方:
    python3 generate_requirements_docx.py requirements.json /mnt/user-data/outputs/要件定義書.docx

章立ては references/document_template.md に従う。
JSONスキーマは requirements_schema.py のdocstringを参照。
"""

import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import requirements_schema as schema  # noqa: E402

JP_FONT = "游ゴシック"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def set_japanese_font(document):
    """既定スタイルに日本語フォントを設定する（東アジア用フォントは別指定が要る）。"""
    style = document.styles["Normal"]
    style.font.name = JP_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), JP_FONT)


def add_heading(document, text, level):
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = JP_FONT
        run.element.rPr.rFonts.set(qn("w:eastAsia"), JP_FONT)
        run.font.color.rgb = ACCENT
    return heading


def add_paragraphs(document, text):
    for block in str(text).split("\n\n"):
        if block.strip():
            document.add_paragraph(block.strip())


def add_table(document, columns, rows, keys):
    """columns=見出し、keys=各列に対応するJSONキー。行が無ければ注記だけ出す。"""
    if not rows:
        note = document.add_paragraph("該当なし（ヒアリング時点で該当項目は挙がっていない）")
        note.runs[0].italic = True
        return

    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cell, header in zip(table.rows[0].cells, columns):
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(9)

    for row in rows:
        cells = table.add_row().cells
        for cell, key in zip(cells, keys):
            cell.text = str(row.get(key, ""))
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def section_intro(document, data):
    project = data["project"]
    add_heading(document, "1. はじめに", 1)

    add_heading(document, "1.1 本書の目的", 2)
    document.add_paragraph(
        f"本書は、{project['name']}において実現すべき要件を定義し、関係者間で合意することを"
        "目的とする。本書で確定するのは「何を実現するか」であり、画面レイアウトや処理方式など"
        "「どう実現するか」は基本設計以降で確定する。"
    )

    add_heading(document, "1.2 対象読者と位置づけ", 2)
    document.add_paragraph(
        "本書は発注者・開発者双方の合意文書であり、以降の基本設計・見積・テスト計画の入力となる。"
        "本書に記載のない事項は原則としてスコープ外とし、追加が必要な場合は変更管理の手続きを経る。"
    )

    if data["glossary"]:
        add_heading(document, "1.3 用語定義", 2)
        add_table(document, ["用語", "定義"], data["glossary"], ["term", "definition"])

    if data["sources"]:
        add_heading(document, "1.4 参照した情報源", 2)
        for source in data["sources"]:
            document.add_paragraph(str(source), style="List Bullet")


def section_background(document, data):
    add_heading(document, "2. システム化の背景と目的", 1)
    purpose = data["project"]["purpose"]
    if purpose:
        add_paragraphs(document, purpose)
    else:
        document.add_paragraph(
            "（未記載）背景と目的が確定していない場合、優先度判断とスコープ調整の拠り所が"
            "無くなるため、確認事項として整理すること。"
        )


def section_business(document, data):
    add_heading(document, "3. 業務要件", 1)
    add_table(
        document,
        ["ID", "業務分類", "業務要件", "詳細", "アクター", "優先度", "状態"],
        data["business_requirements"],
        ["id", "category", "title", "detail", "actor", "priority", "status"],
    )


def section_functional(document, data):
    add_heading(document, "4. 機能要件", 1)

    add_heading(document, "4.1 機能一覧", 2)
    add_table(
        document,
        ["ID", "大分類", "中分類", "機能名", "アクター", "関連業務要件", "優先度", "状態"],
        data["functional_requirements"],
        ["id", "major", "minor", "name", "actor", "related_br", "priority", "status"],
    )

    must = [f for f in data["functional_requirements"] if f.get("priority") == "Must"]
    if must:
        add_heading(document, "4.2 機能詳細（Must）", 2)
        for func in must:
            add_heading(document, f"{func.get('id', '')} {func.get('name', '')}", 3)
            document.add_paragraph(str(func.get("detail", "")))
            meta = document.add_paragraph(
                f"アクター: {func.get('actor', '未確定')} ／ "
                f"関連業務要件: {func.get('related_br', '—')} ／ "
                f"根拠: {func.get('source', '—')}"
            )
            meta.runs[0].font.size = Pt(9)
            meta.runs[0].italic = True


def section_nfr(document, data):
    add_heading(document, "5. 非機能要件", 1)
    document.add_paragraph(
        "各要件は数値と、その数値が成立する条件をセットで記載する。未確定の項目も行として残し、"
        "対応する確認事項のIDを併記する（行が無い項目は見落としに気づけないため削除しない）。"
    )
    grouped = schema.nfr_by_category(data)
    for index, category in enumerate(schema.NFR_CATEGORIES, start=1):
        add_heading(document, f"5.{index} {category}", 2)
        add_table(
            document,
            ["ID", "確認項目", "要件値（数値と条件）", "根拠・背景", "状態", "関連確認事項"],
            grouped.get(category, []),
            ["id", "item", "value", "basis", "status", "question_id"],
        )


def section_constraints(document, data):
    add_heading(document, "6. 制約・前提条件", 1)
    document.add_paragraph(
        "「制約」は動かせない条件、「前提」は現時点でそう仮定している条件を指す。"
        "前提が崩れた場合は再見積・再スケジュールの検討対象となる。"
    )
    add_table(
        document,
        ["ID", "区分", "内容", "崩れた場合の影響", "根拠"],
        data["constraints"],
        ["id", "type", "detail", "impact", "source"],
    )


def section_out_of_scope(document, data):
    add_heading(document, "7. スコープ外", 1)
    document.add_paragraph(
        "本プロジェクトで実施しない事項を示す。備考が「未確認」の行は、"
        "実施しないことの合意が取れていないため、確認事項として回答を得る必要がある。"
    )
    add_table(
        document,
        ["ID", "対象", "理由", "備考"],
        data["out_of_scope"],
        ["id", "detail", "reason", "note"],
    )


def section_questions(document, data):
    add_heading(document, "8. 確認事項（未確定事項）", 1)
    open_questions = [q for q in data["questions"] if q.get("status") != "回答済"]
    document.add_paragraph(
        f"未回答 {len(open_questions)}件。これらが確定するまで、関連する要件は確定扱いにできない。"
    )
    add_table(
        document,
        ["ID", "確認先", "確認内容", "背景・理由", "関連要件", "期限", "状態"],
        data["questions"],
        ["id", "target", "detail", "why", "related", "due", "status"],
    )


def section_revisions(document, data):
    add_heading(document, "9. 改訂履歴", 1)
    revisions = data["revisions"] or [{
        "version": data["project"]["version"] or "0.1",
        "date": data["project"]["date"],
        "note": "初版",
        "author": data["project"]["author"],
    }]
    add_table(
        document,
        ["版数", "日付", "改訂内容", "作成者", "承認者"],
        revisions,
        ["version", "date", "note", "author", "approver"],
    )


def build_cover(document, data):
    project = data["project"]
    title = document.add_heading("要件定義書", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = JP_FONT
        run.element.rPr.rFonts.set(qn("w:eastAsia"), JP_FONT)

    for text in (
        project["name"],
        project["client"],
        f"版数: {project['version']}　作成日: {project['date']}　作成者: {project['author']}",
    ):
        if text:
            paragraph = document.add_paragraph(text)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_page_break()


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)

    data = schema.load(sys.argv[1])
    out_path = schema.prepare_output(sys.argv[2])

    document = Document()
    set_japanese_font(document)

    build_cover(document, data)
    section_intro(document, data)
    section_background(document, data)
    section_business(document, data)
    section_functional(document, data)
    section_nfr(document, data)
    section_constraints(document, data)
    section_out_of_scope(document, data)
    section_questions(document, data)
    section_revisions(document, data)

    document.save(out_path)
    counts = schema.unresolved_counts(data)
    print(f"生成しました: {out_path}")
    print(
        f"未確定の非機能要件 {counts['nfr_unresolved']}件 / "
        f"未回答の確認事項 {counts['questions_open']}件 / "
        f"仮置きの要件 {counts['provisional']}件"
    )


if __name__ == "__main__":
    main()
